from __future__ import annotations

from collections.abc import Callable
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from fileplotter.models import ColumnMetadata, DataSource


ProgressCallback = Callable[[int, str], None]


class ImportFailure(ValueError):
    """An actionable source-import error."""


class DataSourceLoader:
    def workbook_sheets(self, path: str | Path) -> list[str]:
        source_path = Path(path)
        try:
            workbook = load_workbook(source_path, read_only=True, data_only=True)
            try:
                return list(workbook.sheetnames)
            finally:
                workbook.close()
        except Exception as exc:  # openpyxl exposes several format-specific errors
            raise ImportFailure(f"Could not inspect workbook {source_path.name}: {exc}") from exc

    def load(
        self,
        path: str | Path,
        *,
        sheet: str | None = None,
        header_row: int = 0,
        source_id: str | None = None,
        progress: ProgressCallback | None = None,
    ) -> DataSource:
        source_path = Path(path).expanduser().resolve()
        if not source_path.is_file():
            raise ImportFailure(f"Source file does not exist: {source_path}")
        if header_row < 0:
            raise ImportFailure("Header row must be zero or greater.")
        file_format = source_path.suffix.lower().lstrip(".")
        if file_format not in {"csv", "xlsx"}:
            raise ImportFailure("Only comma-delimited CSV and XLSX files are supported.")

        self._progress(progress, 5, f"Opening {source_path.name}")
        try:
            if file_format == "csv":
                frame = pd.read_csv(
                    source_path,
                    header=None,
                    dtype=object,
                    keep_default_na=False,
                    on_bad_lines="error",
                )
                sheet = None
            else:
                available_sheets = self.workbook_sheets(source_path)
                if not available_sheets:
                    raise ImportFailure("The workbook does not contain any worksheets.")
                selected_sheet = sheet or available_sheets[0]
                if selected_sheet not in available_sheets:
                    raise ImportFailure(
                        f"Worksheet {selected_sheet!r} is not present in {source_path.name}."
                    )
                sheet = selected_sheet
                frame = pd.read_excel(
                    source_path,
                    sheet_name=sheet,
                    header=None,
                    dtype=object,
                    engine="openpyxl",
                )
        except ImportFailure:
            raise
        except Exception as exc:
            raise ImportFailure(f"Could not parse {source_path.name}: {exc}") from exc

        self._progress(progress, 45, "Validating columns")
        if frame.empty or header_row >= len(frame.index):
            raise ImportFailure(
                f"Header row {header_row + 1} is outside the file ({len(frame.index)} rows)."
            )
        header = frame.iloc[header_row]
        data = frame.iloc[header_row + 1 :].reset_index(drop=True)
        columns: list[ColumnMetadata] = []
        arrays: dict[str, np.ndarray[Any, np.dtype[np.float64]]] = {}

        for position in range(len(frame.columns)):
            raw_name = header.iloc[position]
            name = str(raw_name).strip() if not pd.isna(raw_name) else ""
            if not name:
                name = f"Column {position + 1}"
            column_id = f"col_{position:04d}"
            raw = data.iloc[:, position]
            text = raw.map(lambda value: "" if pd.isna(value) else str(value).strip())
            nonempty = text.ne("")
            numeric = pd.to_numeric(raw.where(nonempty), errors="coerce").astype(float)
            values = numeric.to_numpy(dtype=np.float64, copy=True)
            infinite = np.isinf(values)
            invalid = int((nonempty & numeric.isna()).sum()) + int(infinite.sum())
            values[infinite] = np.nan
            missing = int((~nonempty).sum())
            numeric_count = int(np.isfinite(values).sum())
            nonempty_count = int(nonempty.sum())
            is_numeric = numeric_count > 0 and (
                nonempty_count == 0 or numeric_count / nonempty_count >= 0.5
            )
            columns.append(
                ColumnMetadata(
                    column_id=column_id,
                    name=name,
                    position=position,
                    numeric=is_numeric,
                    missing_count=missing,
                    invalid_count=invalid,
                )
            )
            if is_numeric:
                values.setflags(write=False)
                arrays[column_id] = values

        if not arrays:
            raise ImportFailure("No numeric columns were found below the selected header row.")
        stat = source_path.stat()
        self._progress(progress, 100, f"Loaded {len(data.index):,} rows")
        kwargs: dict[str, Any] = {
            "path": source_path,
            "format": file_format,
            "sheet": sheet,
            "header_row": header_row,
            "columns": columns,
            "arrays": arrays,
            "fingerprint": {"size": stat.st_size, "mtime_ns": stat.st_mtime_ns},
        }
        if source_id:
            kwargs["source_id"] = source_id
        return DataSource(**kwargs)

    @staticmethod
    def _progress(callback: ProgressCallback | None, value: int, message: str) -> None:
        if callback:
            callback(value, message)
